"""시험표 엑셀 읽기·쓰기 (4장, 9장).

시험항목 시트: ID | 영역 | 시험 항목 | 판정 기준 | 결과 | 수행일 | DUT SW | 수행자 | 비고 | 증거
요약 시트:     영역별 상태 개수 (매번 다시 만든다)
실행이력 시트: 실행할 때마다 한 줄씩 쌓인다

열은 이름이 아니라 위치(A~J)로 읽는다. 그래서 읽고 쓰기 전에 첫 줄의 머리글이 HEADERS 와
같은지 보고, 다르면 멈춘다(열을 끼워 넣거나 옮긴 시험표에 결과를 엉뚱한 칸에 쓰지 않게).

시험표가 원본이고 코드는 시험표를 따른다. 사람이 결과 칸에 N/A 를 적은 항목은
실행하지 않고, 실행기는 N/A 와 INVESTIGATING 을 덮어쓰지 않는다.
"""

from collections import Counter
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .case import STATUSES

SHEET, SUMMARY, HISTORY = "시험항목", "요약", "실행이력"
HEADERS = ["ID", "영역", "시험 항목", "판정 기준", "결과", "수행일", "DUT SW", "수행자", "비고", "증거"]
WIDTHS = [11, 8, 34, 40, 10, 17, 18, 8, 50, 22]
COLORS = {"PASS": "C6EFCE", "FAIL": "FFC7CE", "SKIP": "EDEDED", "BLOCKED": "FFEB9C",
          "N/A": "D9D9D9", "INVESTIGATING": "BDD7EE"}
KEEP = {"N/A", "INVESTIGATING"}  # 사람이 정한 상태
BOLD = Font(bold=True)
HEAD_FILL = PatternFill("solid", fgColor="1F4E78")


def col(name):
    return HEADERS.index(name) + 1


def check_headers(path, first_row):
    got = [str(v).strip() if v is not None else "" for v in list(first_row)[:len(HEADERS)]]
    if got != HEADERS:
        raise SystemExit(f"시험표 머리글이 기준과 다릅니다: {path}\n"
                         f"  기준: {' | '.join(HEADERS)}\n  지금: {' | '.join(got)}\n"
                         "열을 끼워 넣거나 옮기지 말고, 덧붙일 열은 J열(증거) 오른쪽에 두십시오.")


def new_plan(path, rows):
    """rows: [(ID, 영역, 항목, 판정 기준, 결과)] 로 빈 시험표를 만든다."""
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET
    ws.append(HEADERS)
    for r in rows:
        ws.append(list(r) + [""] * (len(HEADERS) - len(r)))
    _style(ws)
    wb.save(path)


def _style(ws):
    for i, w in enumerate(WIDTHS, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = HEAD_FILL
        c.alignment = Alignment(horizontal="center")
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(HEADERS))}{ws.max_row}"
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(vertical="top", wrap_text=True)


def read_plan(path):
    """[{id, area, title, criteria, status}] 를 시험표 순서대로 돌려준다."""
    if not Path(path).exists():
        raise SystemExit(f"시험표가 없습니다: {path}\n먼저 만드십시오: .venv/bin/python plan/make_plan.py")
    ws = load_workbook(path, read_only=True)[SHEET]
    check_headers(path, next(ws.iter_rows(max_row=1, values_only=True), ()))
    items, seen = [], {}
    for n, row in enumerate(ws.iter_rows(min_row=2, values_only=True), 2):
        if not row or not row[0]:
            continue
        tid = str(row[0]).strip()
        if tid in seen:
            raise SystemExit(f"시험표에 같은 ID 가 두 번 있습니다: {tid} ({seen[tid]}행, {n}행)")
        seen[tid] = n
        items.append({"id": str(row[0]).strip(), "area": row[1], "title": row[2],
                      "criteria": row[3], "status": (row[4] or "").strip()})
    return items


def note_for(r):
    parts = [r["reason"]] if r["reason"] else []
    if r["measured"]:
        parts.append(", ".join(f"{k}={v}" for k, v in r["measured"].items()))
    if r["flaky"]:
        first = r["attempts"][0]
        parts.append(f"1차 {first['status']} 후 재시도 PASS ({first['reason'][:60]})")
    if r["console"]:
        levels = Counter(h["level"] for h in r["console"])
        parts.append("콘솔 " + ", ".join(f"{k} {v}건" for k, v in levels.items()))
    if r.get("console_error"):
        parts.append(f"콘솔 확인 못 함({r['console_error'][:60]})")
    return " / ".join(parts)


def apply_results(path, results, meta, run_dir):
    """실행 결과를 시험표에 반영한다. 시험표에 없는 ID 는 돌려준다."""
    wb = load_workbook(path)
    ws = wb[SHEET]
    check_headers(path, [c.value for c in ws[1]])
    rows = {str(ws.cell(r, 1).value).strip(): r for r in range(2, ws.max_row + 1) if ws.cell(r, 1).value}
    unknown = []
    for res in results:
        r = rows.get(res["id"])
        if r is None:
            unknown.append(res["id"])
            continue
        cur = (ws.cell(r, col("결과")).value or "").strip()
        if cur in KEEP:
            continue
        cell = ws.cell(r, col("결과"), res["status"])
        cell.fill = PatternFill("solid", fgColor=COLORS[res["status"]])
        cell.font = BOLD
        ws.cell(r, col("수행일"), res["started"][:16])
        ws.cell(r, col("DUT SW"), meta.get("dut_version", ""))
        ws.cell(r, col("수행자"), meta.get("tester", ""))
        ws.cell(r, col("비고"), note_for(res))
        ws.cell(r, col("증거"), f"{Path(run_dir).name}/{res['log']}")
    _summary(wb, ws)
    _history(wb, results, meta, run_dir)
    tmp = Path(path).with_name(Path(path).name + ".tmp.xlsx")
    wb.save(tmp)  # 저장 도중 끊겨도 원본이 깨지지 않게 임시 파일에 쓰고 바꾼다
    tmp.replace(path)
    return unknown


def _summary(wb, ws):
    if SUMMARY in wb.sheetnames:
        del wb[SUMMARY]
    sm = wb.create_sheet(SUMMARY, 0)
    counts, areas = Counter(), []
    for r in range(2, ws.max_row + 1):
        area = ws.cell(r, col("영역")).value
        if not ws.cell(r, 1).value:
            continue
        status = (ws.cell(r, col("결과")).value or "미수행").strip()
        if area not in areas:
            areas.append(area)
        counts[(area, status)] += 1
    cols = STATUSES + ["미수행"]
    sm.append(["영역"] + cols + ["합계", "통과율"])
    for area in areas + ["전체"]:
        if area == "전체":
            row = [sum(counts[(a, s)] for a in areas) for s in cols]
        else:
            row = [counts[(area, s)] for s in cols]
        total = sum(row)
        run = total - row[cols.index("N/A")] - row[cols.index("미수행")]
        rate = f"{row[0] / run:.0%}" if run else "-"
        sm.append([area] + row + [total, rate])
    for c in sm[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = HEAD_FILL
    for c in sm[sm.max_row]:
        c.font = BOLD
    for i, s in enumerate(cols, 2):
        if s in COLORS:
            sm.cell(1, i).fill = PatternFill("solid", fgColor=COLORS[s])
            sm.cell(1, i).font = BOLD
    sm.column_dimensions["A"].width = 10
    for i in range(2, len(cols) + 4):
        sm.column_dimensions[get_column_letter(i)].width = 15 if i == 7 else 11  # INVESTIGATING


def _history(wb, results, meta, run_dir):
    if HISTORY not in wb.sheetnames:
        h = wb.create_sheet(HISTORY)
        h.append(["실행", "시작", "DUT SW", "수행자", "실행 수", "PASS", "FAIL", "SKIP", "BLOCKED", "흔들림"])
        for c in h[1]:
            c.font = BOLD
        for i, w in enumerate([17, 20, 30, 8, 8, 7, 7, 7, 9, 8], 1):
            h.column_dimensions[get_column_letter(i)].width = w
    h = wb[HISTORY]
    name = Path(run_dir).name
    for row in range(h.max_row, 1, -1):  # 재개한 실행은 같은 줄을 새로 쓴다
        if h.cell(row, 1).value == name:
            h.delete_rows(row)
    c = Counter(r["status"] for r in results)
    h.append([name, meta.get("started", ""), meta.get("dut_version", ""), meta.get("tester", ""),
              len(results), c["PASS"], c["FAIL"], c["SKIP"], c["BLOCKED"], sum(r["flaky"] for r in results)])
