"""실행기 자체 시험. 장비 없이(transport: local) 돌아간다: pytest -q tests"""

import json

import pytest
import yaml

from pitest import case as C
from pitest import logscan, plan, report
from pitest.__main__ import read_exclude
from pitest.lab import load_lab
from pitest.runner import Run


@pytest.fixture
def lab(tmp_path):
    p = tmp_path / "lab.yaml"
    p.write_text(yaml.safe_dump({
        "name": "unit", "tester": "t", "dut_version_cmd": "echo v1",
        "hosts": {"dut": {"transport": "local", "vars": {"lan": "x"}}, "pc": {"transport": "local"}},
    }))
    return load_lab(p)


@pytest.fixture
def registry():
    saved = dict(C.REGISTRY)
    C.REGISTRY.clear()
    yield C.REGISTRY
    C.REGISTRY.clear()
    C.REGISTRY.update(saved)


def run(lab, tmp_path, **kw):
    r = Run(lab, C.select(), root=tmp_path / "results", log=lambda *_: None, **kw)
    return r, {x["id"]: x for x in r.execute()}


def test_statuses_and_evidence(lab, tmp_path, registry):
    C.case("A-1", "A", "pass")(lambda ctx: ctx.dut.run("echo hello", check=True))
    C.case("A-2", "A", "fail")(lambda ctx: C.expect(False, "nope"))
    C.case("A-3", "A", "skip", needs=("lan1",))(lambda ctx: None)
    C.case("A-4", "A", "cmd")(lambda ctx: ctx.dut.run("exit 3", check=True))
    C.case("A-5", "A", "bug")(lambda ctx: 1 / 0)
    r, res = run(lab, tmp_path)
    assert [res[i]["status"] for i in ("A-1", "A-2", "A-3", "A-4", "A-5")] == \
        ["PASS", "FAIL", "SKIP", "BLOCKED", "BLOCKED"]
    assert "ZeroDivisionError" in res["A-5"]["reason"]
    log = (r.dir / "logs" / "A-1.log").read_text()
    assert "$ echo hello" in log and "out| hello" in log and "=> PASS" in log
    assert r.meta["dut_version"] == "v1"


def test_baseline_failure_blocks_rest(lab, tmp_path, registry):
    C.case("B-0", "B", "base", baseline=True)(lambda ctx: C.expect(False, "lab broken"))
    C.case("A-1", "A", "x")(lambda ctx: None)
    _, res = run(lab, tmp_path)
    assert res["A-1"]["status"] == "BLOCKED" and "B-0" in res["A-1"]["reason"]


def test_retry_keeps_first_failure(lab, tmp_path, registry):
    calls = []

    def flaky(ctx):
        calls.append(1)
        C.expect(len(calls) > 1, "first try fails")
    C.case("A-1", "A", "flaky", retry=2)(flaky)
    _, res = run(lab, tmp_path)
    r = res["A-1"]
    assert r["status"] == "PASS" and r["flaky"] and len(r["attempts"]) == 2
    assert r["attempts"][0] == {"status": "FAIL", "reason": "first try fails"}
    assert "1차 FAIL 후 재시도 PASS" in plan.note_for(r)


def test_cleanup_runs_and_failed_cleanup_stops_run(lab, tmp_path, registry):
    done = []

    def t1(ctx):
        ctx.defer(lambda: done.append("restored"))
        C.expect(False, "broken")

    def t2(ctx):
        ctx.defer(lambda: 1 / 0, "restore x")
    C.case("A-1", "A", "x")(t1)
    C.case("A-2", "A", "y")(t2)
    C.case("A-3", "A", "z")(lambda ctx: done.append("ran A-3"))
    r, res = run(lab, tmp_path)
    assert done == ["restored"] and res["A-1"]["status"] == "FAIL"
    assert res["A-2"]["status"] == "BLOCKED" and "복구 실패" in res["A-2"]["reason"]
    assert "A-3" not in res and "LabDirty" in r.state["stopped"]  # 오염된 랩에서 계속 돌지 않는다


def test_cleanup_runs_on_ctrl_c(lab, tmp_path, registry):
    done = []

    def t(ctx):
        ctx.defer(lambda: done.append("restored"))
        raise KeyboardInterrupt
    C.case("A-1", "A", "x")(t)
    r = Run(lab, C.select(), root=tmp_path / "results", log=lambda *_: None)
    with pytest.raises(KeyboardInterrupt):
        r.execute()
    assert done == ["restored"] and r.state["current"] == "A-1" and r.state["stopped"]


def test_resume_and_exclude(lab, tmp_path, registry):
    seen = []
    C.case("B-0", "B", "base", baseline=True)(lambda ctx: seen.append("base"))
    for i in range(3):
        C.case(f"A-{i}", "A", "x")(lambda ctx, i=i: seen.append(i))
    r = Run(lab, C.select(), root=tmp_path / "results", exclude={"A-2": "보드 없음"}, log=lambda *_: None)
    r.state["done"] = ["B-0", "A-0"]  # A-0 까지 하고 끊겼다고 치자
    r.results = [{"id": "B-0", "status": "PASS", "reason": "", "baseline": True},
                 {"id": "A-0", "status": "PASS", "reason": "", "baseline": False}]
    r._save()
    r2 = Run(lab, C.select(), resume=r.dir, exclude={"A-2": "보드 없음"}, log=lambda *_: None)
    results = r2.execute()
    # BASELINE 은 다시 돌고, 끝난 본 시험(A-0)과 제외(A-2)는 돌지 않는다
    assert seen == ["base", 1] and [x["id"] for x in results] == ["A-0", "B-0", "A-1"]
    assert json.loads((r.dir / "run.json").read_text())["state"]["finished"]
    # 다른 조건으로는 이어서 돌릴 수 없다
    with pytest.raises(SystemExit):
        Run(lab, C.select(), resume=r.dir, exclude={}, log=lambda *_: None)


def test_plan_roundtrip_keeps_human_status(lab, tmp_path, registry):
    xlsx = tmp_path / "plan.xlsx"
    plan.new_plan(xlsx, [("A-1", "A", "x", "c"), ("A-2", "A", "y", "c", "N/A"), ("A-3", "B", "z", "c")])
    C.case("A-1", "A", "x")(lambda ctx: None)
    C.case("A-2", "A", "y")(lambda ctx: C.expect(False, "should not overwrite"))
    r, res = run(lab, tmp_path)
    plan.apply_results(xlsx, list(res.values()), r.meta, r.dir)
    items = {i["id"]: i["status"] for i in plan.read_plan(xlsx)}
    assert items == {"A-1": "PASS", "A-2": "N/A", "A-3": ""}
    plan.apply_results(xlsx, list(res.values()), r.meta, r.dir)  # 재개 후 다시 반영해도
    from openpyxl import load_workbook
    assert load_workbook(xlsx)[plan.HISTORY].max_row == 2  # 이력은 한 줄


def test_compare_finds_regression(tmp_path):
    def mk(name, statuses):
        d = tmp_path / name
        d.mkdir()
        (d / "meta.json").write_text("{}")
        (d / "run.json").write_text(json.dumps({"results": [{"id": k, "status": v} for k, v in statuses.items()]}))
        return d
    old = mk("old", {"X": "PASS", "Y": "FAIL", "Z": "SKIP", "W": "PASS"})
    new = mk("new", {"X": "FAIL", "Y": "PASS", "Z": "BLOCKED", "V": "PASS"})
    d = report.compare(old, new)
    assert d["regressed"] == [("X", "PASS", "FAIL")]
    assert d["fixed"] == [("Y", "FAIL", "PASS")]
    assert d["changed"] == [("Z", "SKIP", "BLOCKED")]
    assert report.compare(new, mk("newer", {"X": "FAIL", "Y": "BLOCKED", "Z": "BLOCKED", "V": "PASS"}))["blocked"] \
        == [("Y", "PASS", "BLOCKED")]
    assert d["new"] == [("V", None, "PASS")] and d["missing"] == [("W", "PASS", None)]


def test_logscan_levels_and_allowlist(tmp_path):
    allow = tmp_path / "allow.txt"
    allow.write_text("# comment\nentered forwarding state\n")
    lines = ["br0: port 1 entered forwarding state", "eth0: link error detected",
             "WARNING: CPU: 0 PID: 1 at foo", "Call Trace:", "all good"]
    hits = logscan.scan(lines, logscan.load_allow(allow))
    assert [h["level"] for h in hits] == ["error", "warn", "panic"]


def test_exclude_requires_reason(tmp_path):
    f = tmp_path / "ex.txt"
    f.write_text("A-1  보드가 다른 시험에 쓰임\n# 주석\nA-2\n")
    with pytest.raises(SystemExit):
        read_exclude(f)
    f.write_text("A-1  보드가 다른 시험에 쓰임  # 10/21 까지\n")
    assert read_exclude(f) == {"A-1": "보드가 다른 시험에 쓰임"}


def test_baseline_must_pass_even_if_not_run(lab, tmp_path, registry):
    # BASELINE 결과가 없는 상태(예: 결과 없이 끊김)로는 본 시험을 돌리지 않는다
    C.case("B-0", "B", "base", baseline=True)(lambda ctx: None)
    C.case("A-1", "A", "x")(lambda ctx: None)
    r = Run(lab, [c for c in C.select() if not c.baseline] + [C.REGISTRY["B-0"]],
            root=tmp_path / "results", log=lambda *_: None)
    res = {x["id"]: x for x in r.execute()}
    assert res["A-1"]["status"] == "BLOCKED" and "B-0" in res["A-1"]["reason"]


def test_plan_rejects_duplicate_id_and_missing_file(tmp_path):
    p = tmp_path / "plan.xlsx"
    plan.new_plan(p, [("A-1", "X", "a", "c", ""), ("A-1", "X", "b", "c", "")])
    with pytest.raises(SystemExit, match="같은 ID"):
        plan.read_plan(p)
    with pytest.raises(SystemExit, match="시험표가 없습니다"):
        plan.read_plan(tmp_path / "none.xlsx")


def test_plan_refuses_moved_columns(tmp_path):
    from openpyxl import load_workbook
    p = tmp_path / "plan.xlsx"
    plan.new_plan(p, [("A-1", "X", "a", "c", "")])
    wb = load_workbook(p)
    wb[plan.SHEET].insert_cols(3)  # 열을 하나 끼워 넣으면 위치로 읽는 칸이 모두 어긋난다
    wb.save(p)
    with pytest.raises(SystemExit, match="머리글"):
        plan.read_plan(p)


def test_logscan_uses_journal_priority():
    lines = [json.dumps({"MESSAGE": "hwmon hwmon1: Undervoltage detected!", "PRIORITY": "2"}),
             json.dumps({"MESSAGE": "usb 1-1: device descriptor read failed, retrying", "PRIORITY": "6"}),
             json.dumps({"MESSAGE": "eth0: link down", "PRIORITY": "4"})]
    assert [h["level"] for h in logscan.scan(lines)] == ["crit", "warn"]


def test_dut_version_must_be_readable(tmp_path, registry):
    p = tmp_path / "lab.yaml"
    p.write_text(yaml.safe_dump({"name": "unit", "dut_version_cmd": "false",
                                 "hosts": {"dut": {"transport": "local"}}}))
    C.case("A-1", "A", "x")(lambda ctx: None)
    with pytest.raises(SystemExit, match="DUT 버전을 읽지 못했습니다"):
        Run(load_lab(p), C.select(), root=tmp_path / "results", log=lambda *_: None)


def test_resume_refuses_changed_code(lab, tmp_path, registry, monkeypatch):
    from pitest import runner
    C.case("A-1", "A", "x")(lambda ctx: None)
    r = Run(lab, C.select(), root=tmp_path / "results", log=lambda *_: None)
    r._save()
    monkeypatch.setattr(runner, "code_id", lambda: "changed")
    with pytest.raises(SystemExit, match="시험 코드가 바뀌었습니다"):
        Run(lab, C.select(), resume=r.dir, log=lambda *_: None)


def test_compare_shows_condition_changes_and_measures(tmp_path):
    def mk(name, meta, mbps):
        d = tmp_path / name
        d.mkdir()
        (d / "meta.json").write_text(json.dumps(meta))
        (d / "run.json").write_text(json.dumps({"results": [
            {"id": "PERF-001", "status": "PASS", "measured": {"mbps": mbps}}]}))
        return d
    a = mk("a", {"dut_version": "1", "code": "x"}, 300)
    b = mk("b", {"dut_version": "1", "code": "y"}, 280)
    c = mk("c", {"dut_version": "1", "code": "y"}, 310)
    d = mk("d", {"dut_version": "1", "code": "y"}, 300)
    assert "시험 코드: x -> y" in report.format_compare(report.compare(a, b))
    assert "중앙값 PERF-001   mbps=300" in report.measures([b, c, d], ["PERF-*"])
    assert "중앙값을 내지 않습니다" in report.measures([a, b, c], ["PERF-*"])  # 코드가 다른 실행은 섞지 않는다


def test_spawn_stops_inside_the_host(lab):
    # 일찍 끌 때 장비 안의 프로세스가 신호를 받아 마무리 출력(tcpdump 의 통계 같은)을 남긴다
    p = lab["pc"].spawn("sh -c 'trap \"echo final-stats >&2; exit 0\" TERM; while :; do sleep 0.1; done'", limit=30)
    assert p.pid
    out, err, alive = p.finish(grace=0.2)
    assert alive and "final-stats" in err and "@@pitest-pid" not in err
